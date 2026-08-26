"""Write Maestro Pass/Fail into the ATP Excel Quick Print rows only.

Updates column H (Pass/Fail) and G (Actual Result) on sheet
'New Sprocket Final ATP'. Does not change other modules.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
SUMMARY = REPO / "reports" / "module_runs" / "quick-print_summary.json"
SRC_CANDIDATES = [
    Path(r"C:\Users\HP\Downloads\Hp new sprocket ATP Sheet (1).xlsx"),
    Path(r"C:\Users\HP\Downloads\Hp new sprocket ATP Sheet.xlsx"),
    REPO / "ATP Verification Suite" / "catalog" / "Hp_new_sprocket_ATP_Sheet.xlsx",
]
OUT = REPO / "reports" / "module_runs" / "Hp_Sprocket_ATP_QuickPrint_PassFail.xlsx"
DOWNLOADS = Path(r"C:\Users\HP\Downloads\Hp new sprocket ATP Sheet_QuickPrint_Maestro.xlsx")

EXCEL_TO_FLOW = {
    "QUICK PRINT _001": "QP_001",
    "QUICK PRINT _002": "QP_002",
    "QUICK PRINT_003 (A)": "QP_003A",
    "QUICK PRINT_003 (B)": "QP_003B",
    "QUICK PRINT_003 (C)": "QP_003C",
    "QUICK PRINT_003 (D)": "QP_003D",
    "QUICK PRINT_004": "QP_004",
    "QUICK PRINT_005 (A)": "QP_005A",
    "QUICK PRINT_005 (B)": "QP_005B",
    "QUICK PRINT_006 (A)": "QP_006A",
    "QUICK PRINT_006 (B)": "QP_006B",
    "QUICK PRINT_006 (C)": "QP_006C",
    "QUICK PRINT_006 (D)": "QP_006D",
    "QUICK PRINT_006 (E)": "QP_006E",
    "QUICK PRINT_07": "QP_007",
    "QUICK PRINT_07 (A)": "QP_007A",
    "QUICK PRINT_07 (B)": "QP_007B",
    "QUICK PRINT_08": "QP_008",
    "QUICK PRINT_08 (A)": "QP_008A",
    "QUICK PRINT_08 (B)": "QP_008B",
    "QUICK PRINT_08 (C)": "QP_008C",
    "QUICK PRINT_09": "QP_009",
    "QUICK PRINT_10": "QP_010",
    "QUICK PRINT_11": "QP_011",
    "QUICK PRINT_12": "QP_012",
    "QUICK PRINT_13 (A)": "QP_013A",
    "QUICK PRINT_13 (B)": "QP_013B",
    "QUICK PRINT_13 (C)": "QP_013C",
    "QUICK PRINT_13 (D)": "QP_013D",
    "QUICK PRINT_14": "QP_014",
    "QUICK PRINT_15": "QP_015",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip()).upper()


def _status_by_id(summary: dict) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for r in summary.get("rows") or []:
        m = re.match(r"^(QP_\d+[a-e]?)", r.get("flow") or "", re.I)
        if m:
            out[m.group(1).upper()] = r
    return out


def _match_excel_id(aid: str) -> str | None:
    key = _norm(aid)
    compact = key.replace(" ", "")
    for excel_key, mid in EXCEL_TO_FLOW.items():
        ek = _norm(excel_key)
        if key == ek or compact == ek.replace(" ", ""):
            return mid
    return None


def main() -> int:
    if not SUMMARY.is_file():
        print(f"ERROR: missing {SUMMARY}")
        return 2
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    by_id = _status_by_id(summary)
    src = next((p for p in SRC_CANDIDATES if p.is_file()), None)
    if not src:
        print("ERROR: ATP Excel not found")
        return 2
    wb = load_workbook(src)
    ws = wb["New Sprocket Final ATP"]
    pf = PatternFill("solid", fgColor="C6EFCE")
    ff = PatternFill("solid", fgColor="FFC7CE")
    device = summary.get("device", "")
    updated = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        aid = str(row[0].value or "")
        if "QUICK PRINT" not in aid.upper():
            continue
        mid = _match_excel_id(aid)
        if not mid:
            continue
        rec = by_id.get(mid)
        if not rec:
            continue
        st = rec.get("status", "FAIL")
        reason = (rec.get("failure_reason") or "").strip()
        row[6].value = (
            f"Maestro {st} ({mid}) on {device}"
            + (f": {reason[:180]}" if st != "PASS" and reason else "")
        )
        row[7].value = "Pass" if st == "PASS" else "Fail"
        row[7].fill = pf if st == "PASS" else ff
        row[7].font = Font(bold=True)
        updated += 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    try:
        wb.save(DOWNLOADS)
    except OSError as exc:
        print(f"WARN: could not write Downloads copy: {exc}")
    print(f"updated {updated} Quick Print rows from {src.name}")
    print(f"saved {OUT}")
    print(f"saved {DOWNLOADS}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
