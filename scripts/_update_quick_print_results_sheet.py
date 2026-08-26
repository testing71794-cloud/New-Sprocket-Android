import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

repo = Path(r"d:\Projects-Meastro\New Sprocket Android")
summary_path = repo / "reports/module_runs/quick-print_summary.json"
summary = json.loads(summary_path.read_text(encoding="utf-8"))

full_sec = 9402.9
retry_sec = 1512.5
wall_sec = round(full_sec + retry_sec, 1)

summary["full_suite_execution_time_sec"] = full_sec
summary["retry_failed_execution_time_sec"] = retry_sec
summary["total_wall_clock_sec"] = wall_sec
summary["total_wall_clock"] = (
    f"{int(wall_sec // 3600):02d}:{int((wall_sec % 3600) // 60):02d}:{int(wall_sec % 60):02d}"
)
summary["execution_time_sec"] = wall_sec
summary["run_note"] = (
    "Full suite then --only-failed retry; 4 failures persisted "
    "(QP_013, QP_02, QP_050, QP_17)"
)
summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")

pf = PatternFill("solid", fgColor="C6EFCE")
ff = PatternFill("solid", fgColor="FFC7CE")

xlsx = repo / "reports/module_runs/quick-print_execution_report.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Execution"
ws["A1"] = "HP Sprocket Android - Module quick-print"
ws["A1"].font = Font(bold=True, size=14)
ws.append(["Flow", "Device", "Status", "Failure Reason", "Execution Time (s)", "Timestamp"])
for r in summary["rows"]:
    ws.append(
        [
            r["flow"],
            r.get("device", ""),
            r["status"],
            r.get("failure_reason", ""),
            r["execution_time_sec"],
            r["timestamp"],
        ]
    )
    ws.cell(ws.max_row, 3).fill = pf if r["status"] == "PASS" else ff
sm = wb.create_sheet("Summary")
for k, v in [
    ("module", "quick-print"),
    ("total", summary["total"]),
    ("passed", summary["passed"]),
    ("failed", summary["failed"]),
    ("pass_percent", summary["pass_percent"]),
    ("full_suite_sec", full_sec),
    ("retry_failed_sec", retry_sec),
    ("total_wall_clock_sec", wall_sec),
    ("total_wall_clock", summary["total_wall_clock"]),
    ("device", summary.get("device")),
    ("app", summary.get("app")),
    ("run_finished_utc", datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")),
    ("note", summary["run_note"]),
]:
    sm.append([k, v])
wb.save(xlsx)
print("updated", xlsx)

csv_path = repo / "ATP TestCase Flows/quick-print/atp_quick_print_execution_results.csv"
with csv_path.open("w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(
        [
            "flow",
            "status",
            "execution_time_sec",
            "failure_reason",
            "device",
            "suite_wall_clock",
            "suite_wall_clock_sec",
        ]
    )
    for r in summary["rows"]:
        w.writerow(
            [
                r["flow"],
                r["status"],
                r["execution_time_sec"],
                r.get("failure_reason", ""),
                r.get("device", ""),
                summary["total_wall_clock"],
                wall_sec,
            ]
        )
print("wrote", csv_path)

status_by_id = {}
for r in summary["rows"]:
    m = re.match(r"^(QPX?_\d+)", r["flow"])
    if m:
        status_by_id[m.group(1)] = r["status"]

src = Path(r"C:\Users\HP\Downloads\Hp new sprocket ATP Sheet (1).xlsx")
wb2 = openpyxl.load_workbook(src)
if "Quick Print Maestro" in wb2.sheetnames:
    del wb2["Quick Print Maestro"]
rs = wb2.create_sheet("Quick Print Maestro", 0)
rs["A1"] = "Quick Print Maestro execution results"
rs["A1"].font = Font(bold=True, size=14)
rs.append([])
rs.append(["Device", summary.get("device")])
rs.append(["App", summary.get("app")])
rs.append(["Total flows", summary["total"]])
rs.append(["Passed", summary["passed"]])
rs.append(["Failed", summary["failed"]])
rs.append(["Pass %", summary["pass_percent"]])
rs.append(["Full suite time (sec)", full_sec])
rs.append(["Retry failed time (sec)", retry_sec])
rs.append(["Total wall clock", summary["total_wall_clock"]])
rs.append(["Total wall clock (sec)", wall_sec])
rs.append(
    [
        "Note",
        "Disconnected path: Add Printer → Skip to Connection → "
        "HP Sprocket 200 (DC:A8) → Print → Print Complete",
    ]
)
rs.append(["Failures", "QP_013, QP_02, QP_050, QP_17"])
rs.append([])
rs.append(["Flow", "Status", "Time (s)", "Failure reason"])
for r in summary["rows"]:
    rs.append([r["flow"], r["status"], r["execution_time_sec"], r.get("failure_reason", "")])
    rs.cell(rs.max_row, 2).fill = pf if r["status"] == "PASS" else ff

ws = wb2["New Sprocket Final ATP"]
updated = 0
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    a = row[0].value
    b = row[1].value
    c = row[2].value
    blob = " ".join(str(x or "") for x in [a, b, c])
    if not re.search(r"quick\s*print|select mode|QUICK\s*PRINT", blob, re.I):
        continue
    st = None
    mnum = re.search(r"(?:QUICK\s*PRIN\w*|QPX?)[_\s-]*0*(\d+)", blob, re.I)
    if mnum:
        n = int(mnum.group(1))
        key = f"QPX_{n}" if n <= 29 else None
        if key and key in status_by_id:
            st = status_by_id[key]
    if re.search(
        r"print\s*action|print\s*complete|tap green\s*print|active printing",
        blob,
        re.I,
    ):
        # E2E print path: QP_04 passed this run
        st = "PASS" if status_by_id.get("QP_04") == "PASS" else status_by_id.get("QP_050")
    if not st:
        continue
    row[6].value = f"Maestro {st} on ZA222RFQ75 (wall {summary['total_wall_clock']})"
    row[7].value = "Pass" if st == "PASS" else "Fail"
    comment = str(row[8].value or "").strip()
    if st == "FAIL":
        extra = "See Quick Print Maestro sheet"
        row[8].value = f"{comment} | {extra}".strip(" |") if comment else extra
    row[9].value = "Maestro"
    updated += 1

out = repo / "reports/module_runs/Hp_Sprocket_ATP_with_QuickPrint_Maestro.xlsx"
out2 = Path(r"C:\Users\HP\Downloads\Hp new sprocket ATP Sheet_QuickPrint_Maestro.xlsx")
wb2.save(out)
wb2.save(out2)
print("excel updated quick-print-ish rows", updated)
print("saved", out)
print("saved", out2)
print("WALL", summary["total_wall_clock"], wall_sec)
