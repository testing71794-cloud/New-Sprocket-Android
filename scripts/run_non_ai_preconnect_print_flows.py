"""Run non-AI print-path flows with pre-connected printer assumption (skip AI).

Uses Maestro CLI space-separated flags per docs.maestro.dev.
Writes summary JSON and updates ATP Excel Pass/Fail (Android).
"""
from __future__ import annotations

import json
import re
import subprocess
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\CA Global\Downloads\Hp new sprocket ATP Sheet.xlsx")
OUT_DIR = REPO / "reports" / "module_runs" / "non_ai_preconnect_print"
SUMMARY = OUT_DIR / "non_ai_preconnect_print_summary.json"
MAESTRO = Path(r"C:\Tools\maestro-parallel\bin\maestro.bat")


def _detect_device() -> str:
    import os

    try:
        out = subprocess.check_output(["adb", "devices"], text=True, encoding="utf-8", errors="replace")
    except Exception:
        out = ""
    serials = []
    for line in out.splitlines()[1:]:
        parts = line.split()
        if len(parts) >= 2 and parts[1] == "device":
            serials.append(parts[0])
    env = (os.environ.get("ANDROID_SERIAL") or os.environ.get("MAESTRO_DEVICE") or "").strip()
    if env and env in serials:
        return env
    # Prefer known lab serials if currently connected
    for preferred in ("YLQGQOYS8TONEYNZ", "PVFM5L7LTGHUNZA6"):
        if preferred in serials:
            return preferred
    if serials:
        return serials[0]
    # Fall back to env even if offline (caller will fail clearly)
    return env or "YLQGQOYS8TONEYNZ"


DEVICE = _detect_device()

# excel_id -> flow path (AI intentionally omitted)
FLOWS: list[tuple[str, str]] = [
    ("Print Preview_01", "ATP TestCase Flows/printing/PR_01 - BLE printer connection and Quick Print navigation.yaml"),
    ("Print Preview_02", "ATP TestCase Flows/printing/PR_02 - HP600 Wi-Fi connection and Quick Print navigation.yaml"),
    ("Print Preview_03", "ATP TestCase Flows/printing/PR_03 - HP200 2 x 3 preview canvas.yaml"),
    ("Print Preview_04", "ATP TestCase Flows/printing/PR_04 - Original HP Sprocket 2 x 3 preview canvas.yaml"),
    ("Print Preview_05", "ATP TestCase Flows/printing/PR_05 - HP Sprocket Studio 2 x 3 preview canvas.yaml"),
    ("Print Preview_06", "ATP TestCase Flows/printing/PR_06a - HP Sprocket Studio Plus 2 x 3 preview canvas.yaml"),
    ("Print Preview_07", "ATP TestCase Flows/printing/PR_07 - HP Sprocket Plus 2.3 x 3.4 preview canvas.yaml"),
    ("Print Preview_08", "ATP TestCase Flows/printing/PR_08 - HP Sprocket 3x4 3.5 x 4.25 preview canvas.yaml"),
    ("QUICK PRINT_006 (a)", "ATP TestCase Flows/quick-print/QP_006a - Select Mode user interface.yaml"),
    ("QUICK PRINT_006 (b)", "ATP TestCase Flows/quick-print/QP_006b - Enter Select Mode one photo and print.yaml"),
    ("Photo ID_11", "ATP TestCase Flows/photo-id/PID_11 - Print.yaml"),
    ("Precut_14", "ATP TestCase Flows/precut/PC_14 - Verify print starts on compatible Sprocket 2x3.yaml"),
]


def _extract_reason(blob: str, exit_code: int) -> str:
    # Ignore JAVA_TOOL_OPTIONS noise from stderr
    lines = [
        ln
        for ln in (blob or "").splitlines()
        if ln.strip() and "JAVA_TOOL_OPTIONS" not in ln and "Picked up" not in ln
    ]
    cleaned = "\n".join(lines)
    m = re.search(
        r"(Unable to launch app[^\n]*|Assertion is false[^\n]*|Element not found[^\n]*|Timeout[^\n]*)",
        cleaned,
        re.I,
    )
    if m:
        return m.group(1)[:300]
    if lines:
        return lines[-1][:300]
    return f"exit {exit_code}"


def _run_one(excel_id: str, rel: str, *, max_attempts: int = 2) -> dict:
    flow = REPO / rel
    safe = re.sub(r"[^\w\-]+", "_", excel_id)
    junit = OUT_DIR / f"{safe}.xml"
    debug = OUT_DIR / f"{safe}_debug"
    debug.mkdir(parents=True, exist_ok=True)
    # Global --device before test; space-separated test options.
    # See: https://docs.maestro.dev/maestro-cli/maestro-cli-commands-and-options
    cmd = [
        str(MAESTRO),
        "--no-ansi",
        "--device",
        DEVICE,
        "test",
        "--format",
        "junit",
        "--output",
        str(junit),
        "--debug-output",
        str(debug),
        str(flow),
    ]
    print(f"\n=== RUN {excel_id} ===", flush=True)
    print(" ".join(cmd), flush=True)
    t0 = time.time()
    env = {
        **{k: v for k, v in __import__("os").environ.items()},
        "ANDROID_SERIAL": DEVICE,
        "MAESTRO_CLI_NO_ANSI": "1",
        "NO_COLOR": "1",
        "TERM": "dumb",
    }
    prev = env.get("JAVA_TOOL_OPTIONS", "").strip()
    jansi = (
        "-Djava.net.preferIPv4Stack=true "
        "-Dorg.fusesource.jansi.Ansi.disable=true "
        "-Djansi.passthrough=true"
    )
    env["JAVA_TOOL_OPTIONS"] = f"{prev} {jansi}".strip() if prev else jansi

    last: dict | None = None
    for attempt in range(1, max_attempts + 1):
        log = OUT_DIR / f"{safe}.log"
        try:
            with log.open("w", encoding="utf-8", errors="replace") as lf:
                proc = subprocess.run(
                    cmd,
                    cwd=str(REPO),
                    stdout=lf,
                    stderr=subprocess.STDOUT,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=1200,
                    env=env,
                )
        except subprocess.TimeoutExpired as exc:
            elapsed = round(time.time() - t0, 1)
            log.write_text((log.read_text(encoding="utf-8", errors="replace") if log.exists() else "") + "\nTIMEOUT\n" + str(exc), encoding="utf-8")
            last = {
                "excel_id": excel_id,
                "flow": rel,
                "status": "FAIL",
                "exit_code": -1,
                "elapsed_s": elapsed,
                "failure_reason": "Timeout after 1200s",
                "junit": str(junit),
                "log": str(log),
                "attempts": attempt,
            }
            break
        elapsed = round(time.time() - t0, 1)
        blob = log.read_text(encoding="utf-8", errors="replace") if log.exists() else ""
        status = "PASS" if proc.returncode == 0 else "FAIL"
        reason = _extract_reason(blob, proc.returncode) if status == "FAIL" else ""
        last = {
            "excel_id": excel_id,
            "flow": rel,
            "status": status,
            "exit_code": proc.returncode,
            "elapsed_s": elapsed,
            "failure_reason": reason,
            "junit": str(junit),
            "log": str(log),
            "attempts": attempt,
        }
        print(f"=== {excel_id} {status} ({elapsed}s) attempt {attempt} ===", flush=True)
        if status == "PASS":
            return last
        # Retry only launch/driver flakes
        if attempt < max_attempts and re.search(
            r"Unable to launch app|Maestro driver|UiAutomator|INSTRUMENTATION|not connected|device offline",
            reason + "\n" + blob,
            re.I,
        ):
            print(f"--- retry {excel_id} after launch/driver/device flake ---", flush=True)
            subprocess.run(["adb", "kill-server"], capture_output=True, text=True)
            subprocess.run(["adb", "start-server"], capture_output=True, text=True)
            subprocess.run(
                ["adb", "-s", DEVICE, "wait-for-device"],
                capture_output=True,
                text=True,
                timeout=60,
            )
            subprocess.run(
                ["adb", "-s", DEVICE, "shell", "input", "keyevent", "224"],
                capture_output=True,
                text=True,
            )
            subprocess.run(
                ["adb", "-s", DEVICE, "shell", "am", "force-stop", "com.caglobal.hp.sprocket"],
                capture_output=True,
                text=True,
            )
            time.sleep(3)
            continue
        break
    assert last is not None
    return last


def _match_row(aid: str, mod: str, excel_id: str) -> bool:
    # Case-insensitive; Print Preview_06 appears twice — map Studio Plus only here
    if aid.strip().casefold() == excel_id.casefold():
        if excel_id.casefold() == "print preview_06" and "select" in mod.lower():
            return False
        return True
    compact_a = re.sub(r"\s+", "", aid).casefold()
    compact_e = re.sub(r"\s+", "", excel_id).casefold()
    return compact_a == compact_e


def _update_excel(rows: list[dict]) -> int:
    wb = load_workbook(XLSX)
    ws = wb["New Sprocket Final ATP"]
    pf = PatternFill("solid", fgColor="C6EFCE")
    ff = PatternFill("solid", fgColor="FFC7CE")
    by_id = {r["excel_id"]: r for r in rows}
    updated = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        aid = str(row[0].value or "").strip()
        mod = str(row[1].value or "")
        if aid.startswith("AI Tools"):
            continue
        for excel_id, rec in by_id.items():
            if not _match_row(aid, mod, excel_id):
                continue
            st = rec["status"]
            reason = (rec.get("failure_reason") or "").strip()
            row[6].value = (
                f"Maestro {st} on {DEVICE} (pre-connected Print CTA)"
                + (f": {reason[:180]}" if st != "PASS" and reason else "")
            )
            row[7].value = "Pass" if st == "PASS" else "Fail"
            row[7].fill = pf if st == "PASS" else ff
            row[7].font = Font(bold=True)
            updated += 1
            break
    wb.save(XLSX)
    return updated


def main() -> int:
    import os
    import sys

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    only = {
        a.strip()
        for a in (os.environ.get("ONLY_EXCEL_IDS") or "").split(",")
        if a.strip()
    }
    # Also allow: python ... --only "Print Preview_01,QUICK PRINT_006 (a)"
    if "--only" in sys.argv:
        i = sys.argv.index("--only")
        if i + 1 < len(sys.argv):
            only |= {x.strip() for x in sys.argv[i + 1].split(",") if x.strip()}
    flows = [(e, r) for e, r in FLOWS if not only or e in only]
    if not flows:
        print("No flows selected", flush=True)
        return 2

    # Merge with prior summary when running a subset
    prior_by: dict[str, dict] = {}
    if SUMMARY.exists() and only:
        try:
            prior = json.loads(SUMMARY.read_text(encoding="utf-8"))
            prior_by = {r["excel_id"]: r for r in prior.get("rows") or []}
        except Exception:
            prior_by = {}

    results: list[dict] = []
    for excel_id, rel in flows:
        results.append(_run_one(excel_id, rel))
        merged = dict(prior_by)
        for r in results:
            merged[r["excel_id"]] = r
        # Keep FLOWS order in summary
        ordered = [merged[e] for e, _ in FLOWS if e in merged]
        SUMMARY.write_text(
            json.dumps(
                {
                    "device": DEVICE,
                    "started": datetime.now(timezone.utc).isoformat(),
                    "assumption": "printer pre-connected; Print CTA only; AI skipped",
                    "rows": ordered,
                },
                indent=2,
            ),
            encoding="utf-8",
        )
    n = _update_excel(results)
    passed = sum(1 for r in results if r["status"] == "PASS")
    print(f"\nDone: {passed}/{len(results)} PASS; Excel rows updated: {n}", flush=True)
    print(f"Summary: {SUMMARY}", flush=True)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
