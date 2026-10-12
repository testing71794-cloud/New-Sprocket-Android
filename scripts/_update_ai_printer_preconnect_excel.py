"""Update AI Tools 70/72 conditions for lab pre-connected printer.

Also reinforces Comment notes on physical-print AI rows.
Writes into the user's ATP Excel and PROVIDED_TESTCASES.json.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

REPO = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\CA Global\Downloads\Hp new sprocket ATP Sheet.xlsx")
PROVIDED = REPO / "ATP TestCase Flows" / "ai" / "PROVIDED_TESTCASES.json"

UPDATES = {
    "AI Tools 70": {
        "description": (
            "Verify that Print is displayed and usable when the printer is "
            "pre-connected (no Add Printer / search flow)."
        ),
        "steps": (
            "1. Install and launch the application.\n"
            '2. Verify the Splash Screen displays the HP logo (hp sprocket) with "sprocket" text.\n'
            "3. Wait for the Splash Screen to complete.\n"
            "4. Swipe through the onboarding screens.\n"
            '5. Tap "Get Started".\n'
            '6. Tap "Sign Up Later".\n'
            "7. Wait until the Home screen is displayed.\n"
            '8. If the notification permission popup appears, tap "Allow".\n'
            "9. Ensure a compatible printer is already connected to the app (lab pre-connect).\n"
            '10. Navigate to "AI Tools".\n'
            '11. Tap "Replace Background".\n'
            '12. If the Gallery permission popup appears, tap "Allow All".\n'
            "13. Tap any photo from the Gallery Listing screen.\n"
            '14. Tap the "Generate" button.\n'
            "15. Wait until background generation is completed.\n"
            "16. Tap an available background option.\n"
            '17. Tap "Apply".\n'
            "18. Wait until Print Preview is displayed.\n"
            '19. Verify the green CTA shows "Print" (not "Add Printer").\n'
            "20. Do not start printer search or Add Printer flow."
        ),
        "expected": (
            "With a pre-connected printer, Print Preview shows the Print button "
            "instead of Add Printer. No printer search / Add Printer flow is required."
        ),
        "comment": "Updated: printer pre-connected; tap Print directly (no Add Printer / search).",
    },
    "AI Tools 72": {
        "description": (
            "Verify that Print is displayed on Print Preview when the printer is "
            "pre-connected (Add Printer is not shown)."
        ),
        "steps": (
            "1. Install and launch the application.\n"
            '2. Verify the Splash Screen displays the HP logo (hp sprocket) with "sprocket" text.\n'
            "3. Wait for the Splash Screen to complete.\n"
            "4. Swipe through the onboarding screens.\n"
            '5. Tap "Get Started".\n'
            '6. Tap "Sign Up Later".\n'
            "7. Wait until the Home screen is displayed.\n"
            '8. If the notification permission popup appears, tap "Allow".\n'
            "9. Ensure a compatible printer is already connected to the app (lab pre-connect).\n"
            '10. Navigate to "AI Tools".\n'
            '11. Tap "Replace Background".\n'
            '12. If the Gallery permission popup appears, tap "Allow All".\n'
            "13. Tap any photo from the Gallery Listing screen.\n"
            '14. Tap the "Generate" button.\n'
            "15. Wait until background generation is completed.\n"
            "16. Tap an available background option.\n"
            '17. Tap "Apply".\n'
            "18. Wait until Print Preview is displayed.\n"
            "19. Verify Print is displayed and Add Printer is not displayed."
        ),
        "expected": (
            "The Print button is displayed when the printer is pre-connected. "
            "Add Printer is not shown."
        ),
        "comment": "Updated: printer pre-connected; expect Print CTA (not Add Printer).",
    },
}

PRINT_NOTE_IDS = {
    "AI Tools 14",
    "AI Tools 35",
    "AI Tools 71",
    "AI Tools 73",
    "AI Tools 86",
    "AI Tools 89",
    "AI Tools 93",
    "AI Tools 96",
    "AI Tools 100",
}
NOTE = "Lab: printer pre-connected; Maestro taps Print directly (no Add Printer / search)."


def main() -> int:
    data = json.loads(PROVIDED.read_text(encoding="utf-8"))
    for item in data:
        eid = item.get("excel_id")
        mod = item.get("module") or ""
        if eid == "AI Tools 70":
            item.update({k: v for k, v in UPDATES["AI Tools 70"].items() if k != "comment"})
        elif (
            eid == "AI Tools 72"
            and "Replace Background" in mod
            and "Printer State" in mod
        ):
            item.update({k: v for k, v in UPDATES["AI Tools 72"].items() if k != "comment"})
    PROVIDED.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    wb = load_workbook(XLSX)
    ws = wb["New Sprocket Final ATP"]
    fill = PatternFill("solid", fgColor="FFF2CC")
    updated = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        aid = str(row[0].value or "").strip()
        mod = str(row[1].value or "")
        if aid == "AI Tools 70":
            u = UPDATES["AI Tools 70"]
            row[2].value = u["description"]
            row[3].value = u["steps"]
            row[5].value = u["expected"]
            row[10].value = u["comment"]
            row[10].fill = fill
            updated += 1
        elif aid == "AI Tools 72" and "Replace Background" in mod:
            u = UPDATES["AI Tools 72"]
            row[2].value = u["description"]
            row[3].value = u["steps"]
            row[5].value = u["expected"]
            row[10].value = u["comment"]
            row[10].fill = fill
            updated += 1
        elif aid in PRINT_NOTE_IDS:
            existing = str(row[10].value or "")
            if "pre-connected" not in existing.lower():
                row[10].value = (existing + " | " if existing else "") + NOTE
                row[10].fill = fill
                updated += 1

    wb.save(XLSX)
    print(f"Updated {updated} Excel rows + PROVIDED_TESTCASES.json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
