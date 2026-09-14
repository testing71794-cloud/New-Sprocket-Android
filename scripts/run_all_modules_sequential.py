"""Run every ATP TestCase Flows module; continue after FAIL/CRASH.

Each module uses scripts/run_atp_module_verify.py (recover + continue + Excel).
No extra Maestro YAML: tests keep their existing launch/onboarding/navigation subflows.

Usage:
  python scripts/run_all_modules_sequential.py <device> [maestro] [start_module]
  python scripts/run_all_modules_sequential.py ZA222RFQ75
  python scripts/run_all_modules_sequential.py 16091FDD4004N6 --modules collage,photo-id,home
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from execution.atp_folder_paths import list_atp_modules  # noqa: E402

OUT = REPO / "reports" / "module_runs"


def _parse_args(argv: list[str]) -> argparse.Namespace:
    ap = argparse.ArgumentParser(
        description="Run all ATP modules with shared recovery; continue after failures."
    )
    ap.add_argument("device", nargs="?", default="ZA222RFQ75")
    ap.add_argument("maestro", nargs="?", default=r"C:\Users\HP\maestro\maestro\bin\maestro.bat")
    ap.add_argument(
        "start",
        nargs="?",
        default="",
        help="Start at this module name (legacy positional); skip earlier modules",
    )
    ap.add_argument(
        "--modules",
        default="",
        help="Comma-separated module folders (default: all ATP TestCase Flows modules except common)",
    )
    ap.add_argument("--timeout", default="300")
    ap.add_argument(
        "--skip-printing",
        action="store_true",
        help="Skip the printing module (legacy all-modules default)",
    )
    return ap.parse_args(argv)


def _module_list(args: argparse.Namespace) -> list[str]:
    all_mods = list_atp_modules(REPO)
    if args.modules:
        wanted = [m.strip() for m in args.modules.replace(";", ",").split(",") if m.strip()]
        missing = [m for m in wanted if m not in all_mods]
        if missing:
            print(f"ERROR: unknown module(s): {', '.join(missing)}")
            print(f"Known: {', '.join(all_mods)}")
            raise SystemExit(2)
        mods = wanted
    else:
        mods = list(all_mods)
        if args.skip_printing:
            mods = [m for m in mods if m != "printing"]
    start = (args.start or "").strip()
    if start:
        if start not in mods:
            print(f"ERROR: start module {start!r} not in run list: {', '.join(mods)}")
            raise SystemExit(2)
        mods = mods[mods.index(start) :]
    return mods


def write_overall_excel(index: list[dict], overall: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    OUT.mkdir(parents=True, exist_ok=True)
    xlsx = OUT / "overall_execution_report.xlsx"
    wb = Workbook()
    sm = wb.active
    sm.title = "Overall"
    sm["A1"] = "HP Sprocket Android — all ATP modules"
    sm["A1"].font = Font(bold=True, size=14)
    for k, v in overall.items():
        if isinstance(v, list):
            v = ", ".join(str(x) for x in v)
        sm.append([k, v])
    mods = wb.create_sheet("Modules")
    mods.append(["Module", "Total", "Passed", "Failed", "Crashed", "Pass %", "Seconds", "Devices"])
    pf = PatternFill("solid", fgColor="C6EFCE")
    ff = PatternFill("solid", fgColor="FFC7CE")
    for row in index:
        crashed = int(row.get("crashed") or 0)
        failed = int(row.get("failed") or 0)
        mods.append(
            [
                row.get("module", ""),
                row.get("total", 0),
                row.get("passed", 0),
                failed,
                crashed,
                row.get("pass_percent", 0),
                row.get("execution_time_sec", 0),
                ",".join(row.get("devices") or [])
                if isinstance(row.get("devices"), list)
                else row.get("device", ""),
            ]
        )
        cell = mods.cell(mods.max_row, 4)
        cell.fill = pf if (failed + crashed) == 0 else ff
    flows = wb.create_sheet("All Flows")
    flows.append(["Module", "Id", "Flow", "Device", "Status", "Failure Reason", "Seconds", "Timestamp"])
    fills = {
        "PASS": PatternFill("solid", fgColor="C6EFCE"),
        "FAIL": PatternFill("solid", fgColor="FFC7CE"),
        "CRASH": PatternFill("solid", fgColor="FFEB9C"),
    }
    for row in index:
        mod = row.get("module", "")
        js = OUT / f"{mod}_summary.json"
        if not js.is_file():
            continue
        try:
            data = json.loads(js.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        for r in data.get("rows") or []:
            flows.append(
                [
                    mod,
                    r.get("id") or "",
                    r.get("flow", ""),
                    r.get("device", ""),
                    r.get("status", ""),
                    r.get("failure_reason", ""),
                    r.get("execution_time_sec", ""),
                    r.get("timestamp", ""),
                ]
            )
            flows.cell(flows.max_row, 5).fill = fills.get(r.get("status", ""), fills["FAIL"])
    wb.save(xlsx)
    return xlsx


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv if argv is not None else sys.argv[1:])
    mods = _module_list(args)
    print(f"[all-modules] {len(mods)} module(s): {', '.join(mods)}", flush=True)
    print(f"[all-modules] device={args.device} timeout={args.timeout}", flush=True)
    last_rc = 0
    for mod in mods:
        print(f"\n########## START MODULE {mod} ##########\n", flush=True)
        rc = subprocess.call(
            [
                sys.executable,
                "-u",
                str(REPO / "scripts" / "run_atp_module_verify.py"),
                "--module",
                mod,
                "--device",
                args.device,
                "--maestro",
                args.maestro,
                "--timeout",
                str(args.timeout),
            ],
            cwd=str(REPO),
        )
        print(f"\n########## END MODULE {mod} rc={rc} ##########\n", flush=True)
        if rc != 0:
            last_rc = rc
    index_path = OUT / "index.json"
    if index_path.exists():
        data = json.loads(index_path.read_text(encoding="utf-8"))
        overall = {
            "modules_completed": len(data),
            "total_flows": sum(int(x.get("total") or 0) for x in data),
            "passed": sum(int(x.get("passed") or 0) for x in data),
            "failed": sum(int(x.get("failed") or 0) for x in data),
            "crashed": sum(int(x.get("crashed") or 0) for x in data),
        }
        (OUT / "overall.json").write_text(json.dumps(overall, indent=2), encoding="utf-8")
        print("OVERALL", overall, flush=True)
        xlsx = write_overall_excel(data, overall)
        print(f"overall report: {xlsx}", flush=True)
    return last_rc if last_rc else 0


if __name__ == "__main__":
    raise SystemExit(main())
