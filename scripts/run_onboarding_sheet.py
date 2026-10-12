"""Run Excel onboarding (Sp_*) Maestro flows only.

No Sense/printer setup. Updates Excel Pass/Fail (Android).
CLI: https://docs.maestro.dev/maestro-cli/maestro-cli-commands-and-options
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\CA Global\Downloads\Hp new sprocket ATP Sheet.xlsx")
OUT_DIR = REPO / "reports" / "module_runs" / "onboarding"
MAP = REPO / "reports" / "module_runs" / "full_sheet" / "excel_flow_map.json"
SUMMARY = OUT_DIR / "onboarding_summary.json"
MAESTRO = Path(r"C:\Tools\maestro-parallel\bin\maestro.bat")


def _detect_device() -> str:
    try:
        out = subprocess.check_output(["adb", "devices"], text=True, encoding="utf-8", errors="replace")
    except Exception:
        out = ""
    serials = []
    for line in out.splitlines()[1:]:
        parts = line.split()
        if len(parts) >= 2 and parts[1] == "device":
            serials.append(parts[0])
    env = (os.environ.get("ANDROID_SERIAL") or "").strip()
    if env and env in serials:
        return env
    for preferred in ("PVFM5L7LTGHUNZA6", "YLQGQOYS8TONEYNZ"):
        if preferred in serials:
            return preferred
    return serials[0] if serials else (env or "UNKNOWN")


DEVICE = _detect_device()


def _extract_reason(blob: str, exit_code: int) -> str:
    lines = [
        ln
        for ln in (blob or "").splitlines()
        if ln.strip() and "JAVA_TOOL_OPTIONS" not in ln and "Picked up" not in ln
    ]
    cleaned = "\n".join(lines)
    m = re.search(
        r"(Unable to launch app[^\n]*|Assertion is false[^\n]*|Element not found[^\n]*|Timeout[^\n]*|Parsing Failed[^\n]*|UNAVAILABLE[^\n]*|Device [^\n]*not connected[^\n]*)",
        cleaned,
        re.I,
    )
    if m:
        return m.group(1)[:300]
    for ln in reversed(lines):
        if "Failed" in ln or "FAILED" in ln:
            return ln[:300]
    return f"exit {exit_code}"


def _run_one(excel_id: str, rel: str, max_attempts: int = 2) -> dict:
    flow = REPO / rel
    safe = re.sub(r"[^\w\-]+", "_", excel_id)[:80]
    junit = OUT_DIR / "junit" / f"{safe}.xml"
    debug = OUT_DIR / "debug" / safe
    junit.parent.mkdir(parents=True, exist_ok=True)
    debug.mkdir(parents=True, exist_ok=True)
    cmd = [
        str(MAESTRO),
        "--no-ansi",
        "--device",
        DEVICE,
        "test",
        "--format",
        "junit",
        "--output",
        str(junit),
        "--debug-output",
        str(debug),
        "-e",
        "SKIP_SENSE=1",
        str(flow),
    ]
    print(f"\n=== RUN {excel_id} ===", flush=True)
    t0 = time.time()
    env = {
        **os.environ,
        "ANDROID_SERIAL": DEVICE,
        "MAESTRO_CLI_NO_ANSI": "1",
        "NO_COLOR": "1",
        "TERM": "dumb",
    }
    prev = env.get("JAVA_TOOL_OPTIONS", "").strip()
    jansi = (
        "-Djava.net.preferIPv4Stack=true "
        "-Dorg.fusesource.jansi.Ansi.disable=true "
        "-Djansi.passthrough=true"
    )
    env["JAVA_TOOL_OPTIONS"] = f"{prev} {jansi}".strip() if prev else jansi

    last: dict | None = None
    for attempt in range(1, max_attempts + 1):
        log = OUT_DIR / "logs" / f"{safe}.log"
        log.parent.mkdir(parents=True, exist_ok=True)
        try:
            with log.open("w", encoding="utf-8", errors="replace") as lf:
                proc = subprocess.run(
                    cmd,
                    cwd=str(REPO),
                    env=env,
                    stdout=lf,
                    stderr=subprocess.STDOUT,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                )
            code = proc.returncode
            blob = log.read_text(encoding="utf-8", errors="replace")
        except Exception as exc:
            code = 1
            blob = str(exc)
        elapsed = round(time.time() - t0, 1)
        status = "PASS" if code == 0 else "FAIL"
        reason = "" if status == "PASS" else _extract_reason(blob, code)
        last = {
            "excel_id": excel_id,
            "flow": rel.replace("\\", "/"),
            "status": status,
            "elapsed_s": elapsed,
            "failure_reason": reason,
            "attempts": attempt,
            "log": str(log),
        }
        print(f"=== {excel_id} {status} ({elapsed}s) ===", flush=True)
        if status == "PASS":
            break
    return last or {}


def _match_row(aid: str, mod: str, excel_id: str) -> bool:
    return aid.strip() == excel_id.strip()


def _update_excel_one(rec: dict) -> None:
    wb = load_workbook(XLSX)
    ws = wb["New Sprocket Final ATP"]
    pf = PatternFill("solid", fgColor="C6EFCE")
    ff = PatternFill("solid", fgColor="FFC7CE")
    excel_id = rec["excel_id"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        aid = str(row[0].value or "").strip()
        mod = str(row[1].value or "")
        if not _match_row(aid, mod, excel_id):
            continue
        st = rec["status"]
        reason = (rec.get("failure_reason") or "").strip()
        elapsed = rec.get("elapsed_s", "")
        row[6].value = (
            f"Maestro {st} on {DEVICE} (onboarding, {elapsed}s)"
            + (f": {reason[:160]}" if st != "PASS" and reason else "")
        )
        row[7].value = "Pass" if st == "PASS" else "Fail"
        row[7].fill = pf if st == "PASS" else ff
        row[7].font = Font(bold=True)
        break
    wb.save(XLSX)


def _select_flows() -> list[dict]:
    mapping = json.loads(MAP.read_text(encoding="utf-8"))
    matched = mapping.get("matched") or []
    # Excel Onboarding Splash section only (Sp_*)
    flows = [f for f in matched if re.match(r"(?i)^Sp[_\s-]?\d+", f["excel_id"])]
    # Keep sheet order; unique by excel_id
    seen = set()
    out = []
    for f in flows:
        if f["excel_id"] in seen:
            continue
        seen.add(f["excel_id"])
        out.append(f)
    return out


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not MAP.exists():
        print(f"Missing map: {MAP}", flush=True)
        return 2
    flows = _select_flows()
    print(f"Device: {DEVICE}", flush=True)
    print(f"Onboarding cases: {len(flows)} (Sp_* only)", flush=True)

    subprocess.run(["adb", "-s", DEVICE, "shell", "input", "keyevent", "224"], capture_output=True)

    results: list[dict] = []
    for i, item in enumerate(flows, 1):
        eid, rel = item["excel_id"], item["flow"]
        print(f"\n[{i}/{len(flows)}]", flush=True)
        rec = _run_one(eid, rel, max_attempts=2)
        results.append(rec)
        try:
            _update_excel_one(rec)
        except Exception as exc:
            print(f"Excel update warn: {exc}", flush=True)
        SUMMARY.write_text(
            json.dumps(
                {
                    "device": DEVICE,
                    "updated": datetime.now(timezone.utc).isoformat(),
                    "total": len(flows),
                    "completed": len(results),
                    "passed": sum(1 for r in results if r["status"] == "PASS"),
                    "failed": sum(1 for r in results if r["status"] == "FAIL"),
                    "rows": results,
                },
                indent=2,
            ),
            encoding="utf-8",
        )

    passed = sum(1 for r in results if r["status"] == "PASS")
    print(f"\nDone: {passed}/{len(results)} PASS", flush=True)
    for r in results:
        print(f"  {r['status']}: {r['excel_id']} ({r.get('elapsed_s')}s) {r.get('failure_reason','')[:80]}", flush=True)
    print(f"Summary: {SUMMARY}", flush=True)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
