"""Run AI physical-print / printer-CTA flows with pre-connected printer assumption.

Uses Maestro CLI space-separated flags per docs.maestro.dev.
Writes summary JSON and updates ATP Excel Pass/Fail (Android).
"""
from __future__ import annotations

import json
import re
import subprocess
import time
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
XLSX = Path(r"c:\Users\CA Global\Downloads\Hp new sprocket ATP Sheet.xlsx")
OUT_DIR = REPO / "reports" / "module_runs" / "ai_preconnect_print"
SUMMARY = OUT_DIR / "ai_preconnect_print_summary.json"
MAESTRO = Path(r"C:\Tools\maestro-parallel\bin\maestro.bat")
DEVICE = "YLQGQOYS8TONEYNZ"

# excel_id -> flow path relative to REPO
FLOWS: list[tuple[str, str]] = [
    ("AI Tools 14", "ATP TestCase Flows/ai/AI_14 - Verify that the printed output is generated successfully and matches the image, layout, and selected settings displayed in the Print Preview screen.yaml"),
    ("AI Tools 35", "ATP TestCase Flows/ai/AI_35 - Verify that tapping the Print button on the Print Preview screen navigates the user to the Print State screen.yaml"),
    ("AI Tools 70", "ATP TestCase Flows/ai/AI_70 - Replace Background Printer Connection.yaml"),
    ("AI Tools 71", "ATP TestCase Flows/ai/AI_71 - Replace Background Printing.yaml"),
    ("AI Tools 72", "ATP TestCase Flows/ai/AI_72 - Replace Background Add Printer State.yaml"),
    ("AI Tools 73", "ATP TestCase Flows/ai/AI_73_replace_background - Replace Background Print State.yaml"),
    ("AI Tools 86", "ATP TestCase Flows/ai/AI_86_restore_physical_print - Restore Photo physical print.yaml"),
    ("AI Tools 89", "ATP TestCase Flows/ai/AI_89_enhance_physical_print - Enhance Photo physical print.yaml"),
    ("AI Tools 93", "ATP TestCase Flows/ai/AI_93_lighting_physical_print - Fix Lighting physical print.yaml"),
    ("AI Tools 96", "ATP TestCase Flows/ai/AI_96_upscale_physical_print - Upscale Photo physical print.yaml"),
    ("AI Tools 100", "ATP TestCase Flows/ai/AI_100_generate_object_physical_print - Generate Object physical print.yaml"),
]


def _run_one(excel_id: str, rel: str) -> dict:
    flow = REPO / rel
    safe = re.sub(r"[^\w\-]+", "_", excel_id)
    junit = OUT_DIR / f"{safe}.xml"
    debug = OUT_DIR / f"{safe}_debug"
    debug.mkdir(parents=True, exist_ok=True)
    # Global --device before test; space-separated test options (repo Maestro convention).
    # See: https://docs.maestro.dev/maestro-cli/maestro-cli-commands-and-options
    cmd = [
        str(MAESTRO),
        "--no-ansi",
        "--device",
        DEVICE,
        "test",
        "--format",
        "junit",
        "--output",
        str(junit),
        "--debug-output",
        str(debug),
        str(flow),
    ]
    print(f"\n=== RUN {excel_id} ===", flush=True)
    print(" ".join(cmd), flush=True)
    t0 = time.time()
    env = {
        **dict(**{k: v for k, v in __import__("os").environ.items()}),
        "ANDROID_SERIAL": DEVICE,
        "MAESTRO_CLI_NO_ANSI": "1",
        "NO_COLOR": "1",
        "TERM": "dumb",
    }
    prev = env.get("JAVA_TOOL_OPTIONS", "").strip()
    jansi = (
        "-Djava.net.preferIPv4Stack=true "
        "-Dorg.fusesource.jansi.Ansi.disable=true "
        "-Djansi.passthrough=true"
    )
    env["JAVA_TOOL_OPTIONS"] = f"{prev} {jansi}".strip() if prev else jansi
    proc = subprocess.run(
        cmd,
        cwd=str(REPO),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=1200,
        env=env,
    )
    elapsed = round(time.time() - t0, 1)
    log = OUT_DIR / f"{safe}.log"
    log.write_text(
        (proc.stdout or "") + "\n--- STDERR ---\n" + (proc.stderr or ""),
        encoding="utf-8",
    )
    status = "PASS" if proc.returncode == 0 else "FAIL"
    reason = ""
    if status == "FAIL":
        blob = (proc.stdout or "") + "\n" + (proc.stderr or "")
        m = re.search(r"(Assertion is false[^\n]*|Element not found[^\n]*|Timeout[^\n]*)", blob, re.I)
        reason = (m.group(1) if m else blob.strip().splitlines()[-1] if blob.strip() else f"exit {proc.returncode}")[:300]
    print(f"=== {excel_id} {status} ({elapsed}s) ===", flush=True)
    return {
        "excel_id": excel_id,
        "flow": rel,
        "status": status,
        "exit_code": proc.returncode,
        "elapsed_s": elapsed,
        "failure_reason": reason,
        "junit": str(junit),
        "log": str(log),
    }


def _update_excel(rows: list[dict]) -> int:
    wb = load_workbook(XLSX)
    ws = wb["New Sprocket Final ATP"]
    pf = PatternFill("solid", fgColor="C6EFCE")
    ff = PatternFill("solid", fgColor="FFC7CE")
    by_id = {r["excel_id"]: r for r in rows}
    updated = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        aid = str(row[0].value or "").strip()
        mod = str(row[1].value or "")
        if aid not in by_id:
            continue
        # AI Tools 73 appears twice; only update Replace Background Printer State
        if aid == "AI Tools 73" and "Replace Background" not in mod:
            continue
        if aid == "AI Tools 72" and "Replace Background" not in mod:
            continue
        rec = by_id[aid]
        st = rec["status"]
        reason = (rec.get("failure_reason") or "").strip()
        row[6].value = (
            f"Maestro {st} on {DEVICE} (pre-connected Print CTA)"
            + (f": {reason[:180]}" if st != "PASS" and reason else "")
        )
        row[7].value = "Pass" if st == "PASS" else "Fail"
        row[7].fill = pf if st == "PASS" else ff
        row[7].font = Font(bold=True)
        updated += 1
    wb.save(XLSX)
    return updated


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    results: list[dict] = []
    for excel_id, rel in FLOWS:
        results.append(_run_one(excel_id, rel))
        SUMMARY.write_text(
            json.dumps(
                {
                    "device": DEVICE,
                    "started": datetime.now(timezone.utc).isoformat(),
                    "assumption": "printer pre-connected; Print CTA only",
                    "rows": results,
                },
                indent=2,
            ),
            encoding="utf-8",
        )
    n = _update_excel(results)
    passed = sum(1 for r in results if r["status"] == "PASS")
    print(f"\nDone: {passed}/{len(results)} PASS; Excel rows updated: {n}", flush=True)
    print(f"Summary: {SUMMARY}", flush=True)
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
